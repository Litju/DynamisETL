"""Structural discovery of the verified GymAware landmine press archive.

The discovery receipt is derived from the verified ``LP_data.zip`` central
directory and the decoded structured members:

* per-set GymAware CSV exports (rep-level summary indicators only --- the
  archive proves no sample-level trajectory exists);
* the vision workbook ``GA_vision_data.xlsx``, whose metric sheets carry
  trial identity for 275 numbered sets but numeric values only for the first
  populated row set.

Participant display names in the source are personal data. They are used only
to group rows into deterministic pseudonymous subject identities and are never
persisted in a receipt, a metric or any repository artifact.
"""

from __future__ import annotations

import hashlib
import io
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, BinaryIO

import openpyxl

from dynamis.adapters.gymaware_landmine import authorities
from dynamis.adapters.gymaware_landmine.zip_safety import (
    ZipInspection,
    inspect_zip,
    read_member_bytes,
    structured_members,
)

_CSV_NAME = re.compile(r"^(?P<prefix>.*/)(?P<number>\d{3})\.csv$")
_ASCII_FIELD = re.compile(r"(?P<key>[A-Za-z][A-Za-z ()%]*):\s*(?P<value>[^,]*)")
_REP_METRIC_COLUMNS = (
    "mean_velocity",
    "peak_velocity",
    "peak_power",
    "mean_power",
    "peak_force",
    "mean_force",
)

#: Column header prefixes in the GymAware exports (column order varies by file).
_COLUMN_PREFIXES: tuple[tuple[str, str], ...] = (
    ("mean_velocity", "Conc Mean Velocity"),
    ("peak_velocity", "Conc Peak Velocity"),
    ("peak_power", "Conc Peak Power"),
    ("mean_power", "Conc Mean Power"),
    ("peak_force", "Conc Peak Force"),
    ("mean_force", "Conc Mean Force"),
)


class GymAwareSourceError(ValueError):
    """The archive no longer matches the verified provider structure."""


def decode_text(raw: bytes) -> str:
    """Decode a source text member defensively without assuming one encoding."""
    for encoding in ("utf-8-sig", "gb18030", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")  # pragma: no cover - latin-1 never fails


@dataclass(frozen=True, slots=True)
class GymAwareRep:
    rep_number: int
    values: dict[str, float]
    missing_fields: tuple[str, ...] = ()


@dataclass(frozen=True, slots=True)
class GymAwareSet:
    set_number: int
    member_name: str
    set_id: str | None
    exercise: str | None
    bar_weight_kg: float | None
    total_weight_kg: float | None
    sensor: str | None
    source_date: str | None
    source_time: str | None
    reps: tuple[GymAwareRep, ...]
    issues: tuple[str, ...] = ()
    #: Every source line whose first cell is ``Rep`` (valid, malformed or short).
    source_rep_lines: int = 0

    @property
    def rep_rows(self) -> int:
        return len(self.reps)

    @property
    def malformed_rep_rows(self) -> int:
        return max(self.source_rep_lines - len(self.reps), 0)


def parse_numeric(text: str | None) -> float | None:
    if text is None:
        return None
    candidate = text.strip()
    if candidate in {"", "-", "--"}:
        return None
    try:
        return float(candidate)
    except ValueError:
        return None


def parse_load_kg(raw: Any) -> float | None:
    """Normalize the workbook's loose load representation (``20``, ``20kg``, ``20lg``)."""
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        return float(raw)
    match = re.search(r"-?\d+(?:\.\d+)?", str(raw))
    return float(match.group(0)) if match else None


def parse_gymaware_csv(member_name: str, raw: bytes) -> GymAwareSet:
    """Parse one per-set GymAware export without assuming column order or encoding."""
    text = decode_text(raw)
    lines = text.splitlines()
    match = _CSV_NAME.match(member_name)
    if match is None:
        raise GymAwareSourceError(f"{member_name!r} is not a three-digit GymAware set export")
    set_number = int(match.group("number"))
    metadata: dict[str, str] = {}
    header_line: str | None = None
    header_index = -1
    for index, line in enumerate(lines):
        if line.startswith("Row Type"):
            header_line = line
            header_index = index
            break
    if header_line is not None:
        for key, value in _ASCII_FIELD.findall(lines[0] if lines else ""):
            metadata[key.strip().lower()] = value.strip()
    issues: list[str] = []
    source_rep_lines = sum(1 for line in lines if line.split(",", 1)[0].strip() == "Rep")
    if header_line is None:
        return GymAwareSet(
            set_number=set_number,
            member_name=member_name,
            set_id=metadata.get("set id"),
            exercise=metadata.get("exercise"),
            bar_weight_kg=parse_numeric(metadata.get("bar weight (kg)")),
            total_weight_kg=parse_numeric(metadata.get("total weight (kg)")),
            sensor=metadata.get("sensor"),
            source_date=metadata.get("date"),
            source_time=metadata.get("time"),
            reps=(),
            issues=("no Row Type header found",),
            source_rep_lines=source_rep_lines,
        )
    header = [cell.strip() for cell in header_line.split(",")]
    index_by_metric: dict[str, int] = {}
    for metric, prefix in _COLUMN_PREFIXES:
        for position, column in enumerate(header):
            if column.startswith(prefix):
                index_by_metric[metric] = position
                break
    if len(index_by_metric) != len(_COLUMN_PREFIXES):
        issues.append(f"missing metric columns: {set(_REP_METRIC_COLUMNS) - set(index_by_metric)}")
    reps: list[GymAwareRep] = []
    for line in lines[header_index + 1 :]:
        cells = [cell.strip() for cell in line.split(",")]
        if not cells or cells[0] != "Rep":
            continue
        if len(cells) <= max(1, max(index_by_metric.values(), default=0)):
            issues.append(f"short rep row: {line!r}")
            continue
        rep_number = parse_numeric(cells[1])
        if rep_number is None:
            issues.append(f"rep row without a numeric rep number: {line!r}")
            continue
        values: dict[str, float] = {}
        missing: list[str] = []
        for metric, position in index_by_metric.items():
            parsed = parse_numeric(cells[position])
            if parsed is None:
                missing.append(metric)
            else:
                values[metric] = parsed
        reps.append(
            GymAwareRep(rep_number=int(rep_number), values=values, missing_fields=tuple(missing))
        )
    return GymAwareSet(
        set_number=set_number,
        member_name=member_name,
        set_id=metadata.get("set id"),
        exercise=metadata.get("exercise"),
        bar_weight_kg=parse_numeric(metadata.get("bar weight (kg)")),
        total_weight_kg=parse_numeric(metadata.get("total weight (kg)")),
        sensor=metadata.get("sensor"),
        source_date=metadata.get("date"),
        source_time=metadata.get("time"),
        reps=tuple(reps),
        issues=tuple(issues),
        source_rep_lines=source_rep_lines,
    )


@dataclass(frozen=True, slots=True)
class VisionTrialRow:
    inclusion_number: int
    subject_ordinal: int
    activity: str | None
    load_raw: str | None
    load_kg: float | None
    rep_count: int | None
    values: dict[str, tuple[float | None, float | None, float | None]] = field(default_factory=dict)
    issues: tuple[str, ...] = ()


@dataclass(frozen=True, slots=True)
class VisionWorkbook:
    sheet_names: tuple[str, ...]
    header: tuple[str, ...]
    rows: tuple[VisionTrialRow, ...]
    subject_count: int
    activity_counts: dict[str, int]
    populated_value_count: int
    populated_row_count: int
    issues: tuple[str, ...] = ()


def parse_vision_workbook(
    workbook_source: Path | BinaryIO,
    *,
    sheet_to_metric: dict[str, str],
) -> VisionWorkbook:
    """Parse the vision workbook with the same defensive reading used for CSVs.

    The source may be an in-memory binary stream: the workbook contains
    participant display names, so it is never persisted to a temporary file.
    Display names are consumed only to assign stable pseudonymous subject
    ordinals inside this function; they are not retained in the returned object.
    """
    workbook = openpyxl.load_workbook(workbook_source, read_only=True, data_only=True)
    try:
        sheet_names = tuple(workbook.sheetnames)
        issues: list[str] = []
        unknown = sorted(set(sheet_names) - set(sheet_to_metric))
        if unknown:
            issues.append(f"unmapped sheets: {unknown}")
        identity_order: dict[str, int] = {}
        metadata: dict[int, dict[str, Any]] = {}
        values_by_sheet: dict[str, dict[int, tuple[float | None, float | None, float | None]]] = {}
        header: tuple[str, ...] = ()
        populated_values = 0
        for sheet_name in sheet_names:
            sheet = workbook[sheet_name]
            sheet_rows = list(sheet.iter_rows(values_only=True))
            if not sheet_rows:
                issues.append(f"sheet {sheet_name!r} is empty")
                continue
            header = tuple(str(cell) if cell is not None else "" for cell in sheet_rows[0])
            for raw_row in sheet_rows[1:]:
                if raw_row[0] is None:
                    continue
                inclusion = int(str(raw_row[0]).strip())
                display_name = raw_row[1]
                name = str(display_name).strip() if display_name is not None else ""
                if name and name not in identity_order:
                    identity_order[name] = len(identity_order) + 1
                metadata.setdefault(
                    inclusion,
                    {
                        "subject_ordinal": identity_order.get(name, 0),
                        "activity": str(raw_row[2]) if raw_row[2] is not None else None,
                        "load_raw": str(raw_row[3]) if raw_row[3] is not None else None,
                        "rep_count": (
                            int(raw_row[4]) if isinstance(raw_row[4], (int, float)) else None
                        ),
                    },
                )
                cells = list(raw_row[5:8])
                while len(cells) < 3:
                    cells.append(None)
                triple: tuple[float | None, float | None, float | None] = (
                    parse_numeric(None if cells[0] is None else str(cells[0])),
                    parse_numeric(None if cells[1] is None else str(cells[1])),
                    parse_numeric(None if cells[2] is None else str(cells[2])),
                )
                values_by_sheet.setdefault(sheet_name, {})[inclusion] = triple
                populated_values += sum(1 for value in triple if value is not None)
        activity_counts: Counter[str] = Counter()
        parsed_rows: list[VisionTrialRow] = []
        populated_rows = 0
        for inclusion in sorted(metadata):
            values: dict[str, tuple[float | None, float | None, float | None]] = {}
            row_populated = False
            for sheet_name, metric in sheet_to_metric.items():
                triple = values_by_sheet.get(sheet_name, {}).get(inclusion, (None, None, None))
                values[metric] = triple
                if any(value is not None for value in triple):
                    row_populated = True
            if row_populated:
                populated_rows += 1
            activity = metadata[inclusion]["activity"]
            if activity:
                activity_counts[activity] += 1
            parsed_rows.append(
                VisionTrialRow(
                    inclusion_number=inclusion,
                    subject_ordinal=metadata[inclusion]["subject_ordinal"],
                    activity=activity,
                    load_raw=metadata[inclusion]["load_raw"],
                    load_kg=parse_load_kg(metadata[inclusion]["load_raw"]),
                    rep_count=metadata[inclusion]["rep_count"],
                    values=values,
                )
            )
        return VisionWorkbook(
            sheet_names=sheet_names,
            header=header,
            rows=tuple(parsed_rows),
            subject_count=len(identity_order),
            activity_counts=dict(activity_counts),
            populated_value_count=populated_values,
            populated_row_count=populated_rows,
            issues=tuple(issues),
        )
    finally:
        workbook.close()


@dataclass(frozen=True, slots=True)
class GymAwareDiscovery:
    """Structural receipt of the verified GymAware archive."""

    inspection: ZipInspection
    structured_member_hashes: dict[str, str]
    gymaware_sets: tuple[GymAwareSet, ...]
    vision: VisionWorkbook
    set_numbers_present: tuple[int, ...]
    set_numbers_missing: tuple[int, ...]
    header_variants: dict[str, int]
    row_type_counts: dict[str, int]
    encoding_counts: dict[str, int]
    errors: tuple[str, ...] = ()

    @property
    def gymaware_rep_rows(self) -> int:
        return sum(item.rep_rows for item in self.gymaware_sets)

    def to_dict(self) -> dict[str, Any]:
        return {
            "archive": self.inspection.to_dict(),
            "structured_members": [
                {
                    "name": name,
                    "sha256": digest,
                }
                for name, digest in sorted(self.structured_member_hashes.items())
            ],
            "gymaware": {
                "set_count": len(self.gymaware_sets),
                "rep_rows": self.gymaware_rep_rows,
                "set_numbers_present": list(self.set_numbers_present),
                "set_numbers_missing": list(self.set_numbers_missing),
                "header_variants": self.header_variants,
                "row_type_counts": self.row_type_counts,
                "encoding_counts": self.encoding_counts,
                "loads_raw": _load_representation(self.gymaware_sets),
                "sample_level_trajectory_present": False,
                "errors": list(self.errors),
            },
            "vision": {
                "sheet_names": list(self.vision.sheet_names),
                "header": list(self.vision.header),
                "row_count": len(self.vision.rows),
                "subject_count": self.vision.subject_count,
                "activity_counts": self.vision.activity_counts,
                "load_representation_counts": _vision_load_representation(self.vision),
                "populated_value_count": self.vision.populated_value_count,
                "populated_row_count": self.vision.populated_row_count,
                "inclusion_numbers": [row.inclusion_number for row in self.vision.rows],
                "issues": list(self.vision.issues),
            },
        }


def _load_representation(sets: tuple[GymAwareSet, ...]) -> dict[str, int]:
    counter: Counter[str] = Counter()
    for item in sets:
        if item.bar_weight_kg is None:
            counter["<missing>"] += 1
        else:
            counter[f"{item.bar_weight_kg:g}"] += 1
    return dict(sorted(counter.items()))


def _vision_load_representation(vision: VisionWorkbook) -> dict[str, int]:
    counter: Counter[str] = Counter()
    for row in vision.rows:
        counter[row.load_raw if row.load_raw is not None else "<missing>"] += 1
    return dict(sorted(counter.items()))


def _encoding_of(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030", "latin-1"):
        try:
            raw.decode(encoding)
        except UnicodeDecodeError:
            continue
        return encoding
    return "unknown"


def discover_gymaware_landmine(zip_path: Path | str) -> GymAwareDiscovery:
    """Inspect the verified archive and decode only its structured members."""
    inspection = inspect_zip(zip_path)
    structured = structured_members(inspection)
    if not structured:
        raise GymAwareSourceError("the archive exposes no structured measurement members")
    hashes: dict[str, str] = {}
    csv_sets: list[GymAwareSet] = []
    header_variants: Counter[str] = Counter()
    row_type_counts: Counter[str] = Counter()
    encoding_counts: Counter[str] = Counter()
    errors: list[str] = []
    vision_bytes: bytes | None = None
    for member in structured:
        raw = read_member_bytes(zip_path, member.name)
        hashes[member.name] = hashlib.sha256(raw).hexdigest()
        if member.normalized_name == authorities.VISION_WORKBOOK_KEY:
            vision_bytes = raw
            continue
        if not member.normalized_name.startswith(authorities.GYMAWARE_CSV_PREFIX):
            continue
        text = decode_text(raw)
        encoding_counts[_encoding_of(raw)] += 1
        for line in text.splitlines():
            if line.startswith("Row Type"):
                header_variants[line.strip()] += 1
            head = line.split(",", 1)[0].strip()
            if head in {"Rep", "Total", "Average"}:
                row_type_counts[head] += 1
        csv_sets.append(parse_gymaware_csv(member.name, raw))
    if vision_bytes is None:
        raise GymAwareSourceError(
            f"the verified vision workbook {authorities.VISION_WORKBOOK_KEY!r} is missing"
        )
    collisions: dict[int, list[str]] = {}
    for item in csv_sets:
        collisions.setdefault(item.set_number, []).append(item.member_name)
    duplicated = {number: names for number, names in collisions.items() if len(names) > 1}
    if duplicated:
        raise GymAwareSourceError(
            "the archive maps multiple CSV members onto the same set number: "
            + "; ".join(
                f"{number}: {sorted(names)}" for number, names in sorted(duplicated.items())
            )
        )
    vision = parse_vision_workbook(
        io.BytesIO(vision_bytes), sheet_to_metric=dict(authorities.VISION_SHEET_METRICS)
    )
    present = tuple(sorted(item.set_number for item in csv_sets))
    missing = tuple(sorted(set(range(1, max(present) + 1)) - set(present))) if present else ()
    return GymAwareDiscovery(
        inspection=inspection,
        structured_member_hashes=hashes,
        gymaware_sets=tuple(csv_sets),
        vision=vision,
        set_numbers_present=present,
        set_numbers_missing=missing,
        header_variants=dict(header_variants),
        row_type_counts=dict(row_type_counts),
        encoding_counts=dict(encoding_counts),
        errors=tuple(errors),
    )
